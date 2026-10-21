import os
from pathlib import Path
from typing import Dict, List, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from loguru import logger

def _set_cell_background(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

class DeckPublisher:
    """
    Automated Executive Presentation Deck (.docx) Publisher for Bourbon.AI.
    Generates agency-quality campaign proposals complete with:
    - Cover Header & Brand Specifications Table
    - High-Res Visual Plates & Composited Ad Images
    - Ad Copywriting, Headline Hierarchies, and Meta Ad Formats
    - Sign-Off & Client Approval Block
    """

    COLOR_PRIMARY = RGBColor(0, 28, 84)       # Deep Navy (#001C54)
    COLOR_SECONDARY = RGBColor(245, 166, 35)  # Amber Gold (#F5A623)
    COLOR_DARK = RGBColor(30, 30, 30)         # Charcoal Dark (#1E1E1E)
    COLOR_MUTED = RGBColor(100, 110, 120)

    @classmethod
    def build_campaign_deck(
        cls,
        campaign_title: str,
        subtitle: str,
        client_name: str,
        specs: List[tuple],
        slides: List[Dict[str, Any]],
        output_path: Path
    ) -> Path:
        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # 1. Title Banner
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(campaign_title.upper())
        title_run.font.name = "Arial"
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = cls.COLOR_PRIMARY
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(subtitle)
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.color.rgb = cls.COLOR_SECONDARY
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # 2. Campaign Specifications Table
        if specs:
            spec_table = doc.add_table(rows=len(specs), cols=2)
            spec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            spec_table.autofit = False

            for idx, (k, v) in enumerate(specs):
                row = spec_table.rows[idx]
                c0, c1 = row.cells[0], row.cells[1]
                c0.width, c1.width = Inches(2.2), Inches(4.6)
                _set_cell_background(c0, "EBF3FA")
                _set_cell_background(c1, "FAFAFA")
                r0 = c0.paragraphs[0].add_run(str(k))
                r0.font.bold = True
                r0.font.size = Pt(9.5)
                r0.font.color.rgb = cls.COLOR_PRIMARY
                r1 = c1.paragraphs[0].add_run(str(v))
                r1.font.size = Pt(9.5)
                r1.font.color.rgb = cls.COLOR_DARK

            doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 3. Slide Decks
        for s in slides:
            doc.add_page_break()

            # Slide Badge Header
            badge_p = doc.add_paragraph()
            badge_run = badge_p.add_run(s.get("badge", f"SLIDE {s.get('slide_num', 1)}"))
            badge_run.font.bold = True
            badge_run.font.size = Pt(13)
            badge_run.font.color.rgb = cls.COLOR_PRIMARY

            head_p = doc.add_paragraph()
            h_main = head_p.add_run(f'"{s.get("headline", "")}" ')
            h_main.font.bold = True
            h_main.font.size = Pt(15)
            h_main.font.color.rgb = cls.COLOR_DARK

            if s.get("accent"):
                h_acc = head_p.add_run(f'— {s.get("accent")}')
                h_acc.font.bold = True
                h_acc.font.size = Pt(15)
                h_acc.font.color.rgb = cls.COLOR_SECONDARY

            # Insert Image
            img_path = s.get("image_path")
            p_img = None
            if img_path:
                p_cand = Path(img_path)
                if p_cand.exists() and p_cand.is_file():
                    p_img = p_cand
                else:
                    base_dir = Path(__file__).parent.parent
                    p_cand2 = base_dir / str(img_path).lstrip("/\\")
                    if p_cand2.exists() and p_cand2.is_file():
                        p_img = p_cand2

            if p_img:
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(str(p_img), width=Inches(4.8))
                img_p.paragraph_format.space_after = Pt(8)

            # Details Table
            card_rows = [
                ("Skill / Program Focus:", s.get("skill_focus", "")),
                ("Visual Scene Concept:", s.get("visual_desc", "")),
                ("Primary Ad Copy:", s.get("ad_copy", "")),
                ("Direct Call to Action (CTA):", s.get("cta", "Book a Free 1-on-1 Discovery Demo"))
            ]
            t = doc.add_table(rows=len(card_rows), cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False

            for idx, (label, val) in enumerate(card_rows):
                r = t.rows[idx]
                c0, c1 = r.cells[0], r.cells[1]
                c0.width, c1.width = Inches(2.2), Inches(4.6)
                _set_cell_background(c0, "F0F4F8")
                _set_cell_background(c1, "FFFFFF")
                r0 = c0.paragraphs[0].add_run(label)
                r0.font.bold = True
                r0.font.size = Pt(9.0)
                r0.font.color.rgb = cls.COLOR_PRIMARY
                r1 = c1.paragraphs[0].add_run(str(val))
                r1.font.size = Pt(9.0)
                r1.font.color.rgb = cls.COLOR_DARK

        # 4. Sign-Off Page
        doc.add_page_break()
        end_p = doc.add_paragraph()
        end_run = end_p.add_run("CAMPAIGN APPROVAL & SIGN-OFF")
        end_run.font.bold = True
        end_run.font.size = Pt(14)
        end_run.font.color.rgb = cls.COLOR_PRIMARY

        sign_table = doc.add_table(rows=3, cols=2)
        sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sign_table.autofit = False
        sign_rows = [
            ("Client Approval Status:", "APPROVED / READY FOR PRODUCTION"),
            ("Reviewer & Marketing Lead:", f"{client_name} Brand Management"),
            ("Production Delivery:", "1080x1080 High-Res PNGs + Hootsuite JSON Export")
        ]
        for idx, (lbl, val) in enumerate(sign_rows):
            r = sign_table.rows[idx]
            c0, c1 = r.cells[0], r.cells[1]
            c0.width, c1.width = Inches(2.5), Inches(4.3)
            _set_cell_background(c0, "EBF3FA")
            _set_cell_background(c1, "FAFAFA")
            r0 = c0.paragraphs[0].add_run(lbl)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            r0.font.color.rgb = cls.COLOR_PRIMARY
            r1 = c1.paragraphs[0].add_run(val)
            r1.font.size = Pt(9.5)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"Generated Presentation Deck at: {output_path}")
        return output_path
